import base64
import io
import mimetypes
from pathlib import Path
from typing import Optional, Tuple

MAX_FILE = 20 * 1024 * 1024
IMAGE_EXTS = {'.png', '.bmp', '.gif', '.jpg', '.heic', '.jpeg', '.webp', '.tiff'}


def parse_file(fname: str, data: bytes, max_file=MAX_FILE) -> Tuple[str, Optional[str], Optional[str]]:
    if len(data) > max_file:
        raise ValueError("Arquivo maior que 20 MB")
    ext = Path(fname).suffix.lower()
    if ext in IMAGE_EXTS:
        mime = mimetypes.guess_type(fname)[0] or "image/jpeg"
        b64 = base64.b64encode(data).decode()
        return f"[Imagem {fname} - sera analisada por visao]\n", b64, mime

    if ext in {
        ".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".html", ".css",
        ".sh", ".java", ".c", ".cpp", ".rs", ".go", ".yaml", ".yml", ".ini", ".log",
    }:
        try:
            txt = data.decode(errors="ignore")
        except Exception:
            txt = ""
        return txt[:200000], None, None

    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            parts = []
            count = 0
            for page in reader.pages[:40]:
                text = (page.extract_text() or "")[:200000 - count]
                parts.append(text)
                count += len(text)
                if count >= 200000:
                    break
            txt = "\n".join(parts)[:200000]
            if len(txt.strip()) >= 50:
                return txt, None, None
            try:
                import fitz
                with fitz.open(stream=data, filetype="pdf") as doc:
                    if len(doc):
                        pix = doc[0].get_pixmap(dpi=150)
                        return (
                            f"[PDF escaneado {fname} - primeira pagina enviada pra OCR]\n",
                            base64.b64encode(pix.tobytes("png")).decode(), "image/png",
                        )
            except ImportError:
                pass
            return txt, None, None
        except Exception as error:
            raise ValueError(f"Não foi possível ler o PDF: {error}") from error

    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return "\n".join(parts)[:200000], None, None
        except Exception as error:
            raise ValueError(f"Não foi possível ler o DOCX: {error}") from error

    if ext in {".xlsx", ".xls"}:
        try:
            import pandas as pd
            df = pd.read_excel(io.BytesIO(data), nrows=20)
            return (
                f"[Planilha {fname}: amostra de até 20 linhas x {df.shape[1]} colunas]\n"
                f"{df.to_string(index=False)[:200000]}\n", None, None,
            )
        except Exception as error:
            raise ValueError(f"Não foi possível ler a planilha: {error}") from error

    return f"[Arquivo {fname} {len(data)} bytes - tipo nao parseado, apenas nome]", None, None
