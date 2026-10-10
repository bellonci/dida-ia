import base64
import http.client
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
CONFIG = tempfile.TemporaryDirectory(prefix='dida_tests_')
os.environ['DIDA_CONFIG_DIR'] = CONFIG.name
spec = importlib.util.spec_from_file_location('dida', ROOT / 'Ia-Completa-Hibrida.py')
dida = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = dida
spec.loader.exec_module(dida)


class CoreTests(unittest.TestCase):
    def test_equal_length_correlation(self):
        signal = dida.np.sin(dida.np.arange(320) / 7)
        correlation, delay = dida.norm_xcorr(signal, signal.copy())
        self.assertGreater(correlation, 0.99)
        self.assertEqual(delay, 0)

    def test_short_correlation(self):
        self.assertEqual(dida.norm_xcorr(dida.np.ones(10), dida.np.ones(20)), (0.0, 0))

    def test_sensitivity_rejects_invalid_numbers(self):
        for value in [float('nan'), float('inf'), -0.1, 1.1, 'bad']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                dida.validate_sensitivity(value)

    def test_invalid_saved_sensitivity_uses_default(self):
        with patch.object(dida, 'MIC_JSON', Path(CONFIG.name) / 'bad-mic.json'):
            dida.MIC_JSON.write_text('{"sensibilidade": NaN}')
            self.assertEqual(dida.load_sens(), 0.5)

    def test_atomic_private_json_and_cleanup(self):
        path = Path(CONFIG.name) / 'nested' / 'profile.json'
        dida.write_private_json(path, {'name': 'João'})
        self.assertEqual(json.loads(path.read_text()), {'name': 'João'})
        self.assertEqual(path.stat().st_mode & 0o777, 0o600)
        with self.assertRaises(ValueError):
            dida.write_private_json(path, {'value': float('nan')})
        self.assertEqual(json.loads(path.read_text()), {'name': 'João'})
        self.assertEqual(list(path.parent.glob('*.tmp')), [])

    def test_memory_uses_its_own_parent(self):
        path = Path(CONFIG.name) / 'other' / 'memory.json'
        memory = dida.Memory(path)
        memory.add('Olá', 'Oi')
        restored = dida.Memory(path)
        self.assertEqual(restored.data['conversas'][0]['user'], 'Olá')

    def test_corrupt_memory_backup_failure_does_not_crash(self):
        path = Path(CONFIG.name) / 'corrupt.json'
        path.write_text('invalid')
        with patch.object(dida.shutil, 'copy2', side_effect=PermissionError('denied')):
            memory = dida.Memory(path)
        self.assertEqual(memory.data, {'conversas': []})

    def test_docx_includes_tables_without_temp_files(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph('Texto normal')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Valor na tabela'
        data = io.BytesIO()
        doc.save(data)
        with patch.object(dida.tempfile, 'mkstemp', side_effect=AssertionError('no disk')):
            text, image, mime = dida.parse_file('relatorio.docx', data.getvalue())
        self.assertIn('Texto normal', text)
        self.assertIn('Valor na tabela', text)
        self.assertIsNone(image)

    def test_pdf_without_disk(self):
        import pypdf
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=100, height=100)
        data = io.BytesIO()
        writer.write(data)
        text, image, mime = dida.parse_file('blank.pdf', data.getvalue())
        self.assertIsInstance(text, str)

    def test_spreadsheet_is_limited_sample(self):
        import pandas as pd
        data = io.BytesIO()
        pd.DataFrame({'value': list(range(100))}).to_excel(data, index=False)
        text, _, _ = dida.parse_file('test.xlsx', data.getvalue())
        self.assertIn('amostra', text)
        self.assertNotIn('99', text)

    def test_oversized_attachment_rejected(self):
        with patch.object(dida, 'MAX_FILE', 3), self.assertRaises(ValueError):
            dida.parse_file('test.txt', b'abcd')

    def test_latest_message_replaces_pending(self):
        from dida_app.requests import RequestCoordinator
        manager = RequestCoordinator()
        with patch.object(dida, 'requests_manager', manager), patch.object(dida, 'request_queue', manager.queue):
            dida.submit_message('primeira')
            first = manager.current
            dida.submit_message('segunda')
            self.assertTrue(first.cancelled())
            request = manager.queue.get_nowait()
            self.assertEqual(request.message, 'segunda')
            manager.queue.task_done()
            self.assertEqual(manager.queue.unfinished_tasks, 0)

    def test_invalid_message_rejected(self):
        for value in ['', ' ', None, 42]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                dida.submit_message(value)

    def test_local_model_stream_and_response_close(self):
        from unittest.mock import Mock
        response = Mock(status_code=200)
        response.iter_lines.return_value = [b'data: {"choices":[{"delta":{"content":"Oi"}}]}', b'data: [DONE]']
        dida.estado.interrupt_flag.clear()
        with patch.object(dida.model_catalog, 'refresh', return_value=True), patch('dida_app.llm.open_stream', return_value=response) as post:
            self.assertEqual(list(dida.local_chat_stream([])), ['Oi'])
        self.assertEqual(post.call_args.args[0], 'http://127.0.0.1:8081/v1/chat/completions')
        self.assertEqual(post.call_args.args[1]['model'], dida.local_model.model_name)
        response.close.assert_called_once()


class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = dida.ThreadingHTTPServer(('127.0.0.1', 0), dida.Handler)
        cls.server.daemon_threads = True
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()

    def request(self, method, path, body=None, headers=None):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=5)
        try:
            connection.request(method, path, body=body, headers=headers or {})
            response = connection.getresponse()
            return response.status, response.read()
        finally:
            connection.close()

    def test_page(self):
        status, body = self.request('GET', '/')
        self.assertEqual(status, 200)
        self.assertIn(b'<html lang="pt-BR">', body)

    def test_rebinding_host_rejected(self):
        self.assertEqual(self.request('GET', '/', headers={'Host': 'evil.example'})[0], 403)

    def test_cross_origin_mutation_rejected(self):
        self.assertEqual(self.request('POST', '/api/profile', '{}', {'Origin': 'https://evil.example'})[0], 403)

    def test_malformed_origin_rejected(self):
        self.assertEqual(self.request('GET', '/', headers={'Origin': 'http://['})[0], 403)

    def test_same_origin_allowed(self):
        origin = f'http://127.0.0.1:{self.server.server_port}'
        self.assertEqual(self.request('POST', '/api/profile', '{}', {'Origin': origin})[0], 200)

    def test_json_object_validation(self):
        for body in ['[]', 'null', '{', '"text"', b'\xff']:
            with self.subTest(body=body):
                self.assertEqual(self.request('POST', '/api/msg', body)[0], 400)

    def test_invalid_message_http(self):
        for body in ['{"texto":42}', '{"texto":""}']:
            self.assertEqual(self.request('POST', '/api/msg', body)[0], 400)

    def test_sensitivity_nan_rejected(self):
        before = dida.estado.sensibilidade
        self.assertEqual(self.request('POST', '/api/sens', '{"v":NaN}')[0], 400)
        self.assertEqual(dida.estado.sensibilidade, before)

    def test_oversized_body_rejected(self):
        self.assertEqual(self.request('POST', '/api/msg', 'x' * 65537)[0], 413)

    def test_csrf_still_required_for_files(self):
        self.assertEqual(self.request('POST', '/api/files/permission', '{"enabled":true}')[0], 403)

    def test_text_attachment_queued(self):
        body = json.dumps({'fname': 'example.txt', 'data': base64.b64encode(b'Conteudo do anexo').decode()})
        from dida_app.requests import RequestCoordinator
        manager = RequestCoordinator()
        with patch.object(dida, 'requests_manager', manager), patch.object(dida, 'request_queue', manager.queue):
            status, response = self.request('POST', '/anexo', body, {'Content-Type': 'application/json'})
            self.assertEqual(status, 200)
            self.assertIn('Conteudo do anexo', dida.request_queue.get_nowait()[3]['text'])


if __name__ == '__main__':
    unittest.main()
